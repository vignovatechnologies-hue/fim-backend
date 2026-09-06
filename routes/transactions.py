import datetime
import io
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Body
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from database import get_db
from models import Transaction, Budget, User
from schemas import TransactionCreate, TransactionResponse, BudgetCategoryResponse
from dependencies import get_current_user

router = APIRouter(tags=["transactions"])

BUDGET_COLORS = {
    "Food & Dining":   "bg-amber-100 text-amber-700",
    "Shopping":        "bg-rose-100 text-rose-700",
    "Transport":       "bg-sky-100 text-sky-700",
    "Entertainment":   "bg-violet-100 text-violet-700",
    "Home & Bills":    "bg-emerald-100 text-emerald-700",
}

CAT_MAP = {
    "Food":          "Food & Dining",
    "Shopping":      "Shopping",
    "Transport":     "Transport",
    "Entertainment": "Entertainment",
    "Home":          "Home & Bills",
    "Bills":         "Home & Bills",
    "Healthcare":    "Shopping",
    "Income":        "Income",
    "Salary":        "Income",
    "Freelance":     "Income",
    "Investments":   "Income",
    "Gift":          "Income",
    "Other Income":  "Income",
    "Savings":       "Savings",
}

INCOME_CATEGORIES = {"Income", "Salary", "Freelance", "Investments", "Gift", "Other Income"}

def _serialize_txn(txn) -> dict:
    return TransactionResponse.from_orm_model(txn).model_dump()


@router.get("/api/transactions")
def get_transactions(
    month: Optional[int] = None,
    year: Optional[int] = None,
    all_time: Optional[bool] = False,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    query = db.query(Transaction).filter(Transaction.user_id == current_user.id)
    
    if not all_time:
        now = datetime.datetime.utcnow()
        m = month if (month is not None and month > 0) else now.month
        y = year if (year is not None and year > 0) else now.year
        
        start_of_month = datetime.datetime(y, m, 1)
        if m == 12:
            end_of_month = datetime.datetime(y + 1, 1, 1)
        else:
            end_of_month = datetime.datetime(y, m + 1, 1)
        query = query.filter(Transaction.when >= start_of_month, Transaction.when < end_of_month)
        
    txns = query.order_by(Transaction.when.desc()).limit(200).all()
    return JSONResponse([_serialize_txn(t) for t in txns])


# ── ADD a transaction (expense / income) ──────────────────────────────────────
@router.post("/api/transactions")
def add_transaction(
    txn_data: TransactionCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    raw_cat = txn_data.category.strip()
    std_cat = CAT_MAP.get(raw_cat, raw_cat)
    
    # ── Strict Income vs Expense determination ─────────────────────────────
    is_income = (
        (txn_data.payment_status and txn_data.payment_status.lower() == "credit")
        or raw_cat in INCOME_CATEGORIES
        or std_cat == "Income"
        or (txn_data.amount > 0 and (not txn_data.payment_status or txn_data.payment_status.lower() != "debit"))
    )

    if is_income:
        final_amount = abs(txn_data.amount)
        payment_status = "credit"
        final_category = raw_cat if raw_cat in INCOME_CATEGORIES else "Income"
    else:
        final_amount = -abs(txn_data.amount)
        payment_status = "debit"
        final_category = std_cat if std_cat in CAT_MAP.values() else raw_cat

    txn_when = datetime.datetime.utcnow()
    if txn_data.when and txn_data.when.strip():
        try:
            txn_when = datetime.datetime.fromisoformat(txn_data.when.strip().replace("Z", ""))
        except Exception:
            pass

    txn = Transaction(
        user_id=current_user.id,
        name=txn_data.name.strip(),
        category=final_category,
        amount=final_amount,
        payment_status=payment_status,
        when=txn_when,
    )
    db.add(txn)
    db.commit()
    db.refresh(txn)
    return JSONResponse(_serialize_txn(txn))


# ── DELETE a transaction ──────────────────────────────────────────────────────
@router.delete("/api/transactions/{txn_id}")
def delete_transaction(
    txn_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    txn = db.query(Transaction).filter(
        Transaction.id == txn_id,
        Transaction.user_id == current_user.id
    ).first()
    if not txn:
        raise HTTPException(status_code=404, detail="Transaction not found")

    db.delete(txn)
    db.commit()
    return JSONResponse({"status": "success", "message": "Transaction deleted successfully", "id": txn_id})


# ── UPDATE / EDIT a transaction ─────────────────────────────────────────────
@router.put("/api/transactions/{txn_id}")
def update_transaction(
    txn_id: int,
    txn_data: TransactionCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    txn = db.query(Transaction).filter(
        Transaction.id == txn_id,
        Transaction.user_id == current_user.id
    ).first()
    if not txn:
        raise HTTPException(status_code=404, detail="Transaction not found")

    raw_cat = txn_data.category.strip()
    std_cat = CAT_MAP.get(raw_cat, raw_cat)

    is_income = (
        (txn_data.payment_status and txn_data.payment_status.lower() == "credit")
        or raw_cat in INCOME_CATEGORIES
        or std_cat == "Income"
        or (txn_data.amount > 0 and (not txn_data.payment_status or txn_data.payment_status.lower() != "debit"))
    )

    if is_income:
        final_amount = abs(txn_data.amount)
        payment_status = "credit"
        final_category = raw_cat if raw_cat in INCOME_CATEGORIES else "Income"
    else:
        final_amount = -abs(txn_data.amount)
        payment_status = "debit"
        final_category = std_cat if std_cat in CAT_MAP.values() else raw_cat

    if txn_data.when and txn_data.when.strip():
        try:
            txn.when = datetime.datetime.fromisoformat(txn_data.when.strip().replace("Z", ""))
        except Exception:
            pass

    txn.name = txn_data.name.strip()
    txn.category = final_category
    txn.amount = final_amount
    txn.payment_status = payment_status

    db.commit()
    db.refresh(txn)
    return JSONResponse(_serialize_txn(txn))


# ── GET budgets with real spent amounts ───────────────────────────────────────
@router.get("/api/budgets")
def get_budgets(
    month: Optional[int] = None,
    year: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    budgets = db.query(Budget).filter(Budget.user_id == current_user.id).all()

    now = datetime.datetime.utcnow()
    m = month if month is not None else now.month
    y = year if year is not None else now.year

    start_of_month = datetime.datetime(y, m, 1)
    if m == 12:
        end_of_month = datetime.datetime(y + 1, 1, 1)
    else:
        end_of_month = datetime.datetime(y, m + 1, 1)

    result = []
    for b in budgets:
        spent = (
            db.query(Transaction)
            .filter(
                Transaction.user_id == current_user.id,
                Transaction.category == b.category,
                Transaction.when >= start_of_month,
                Transaction.when < end_of_month,
                Transaction.amount < 0,
            )
            .all()
        )
        total_spent = sum(abs(t.amount) for t in spent)
        result.append({
            "name":   b.category,
            "spent":  round(total_spent, 2),
            "budget": round(b.budget_amount, 2),
            "color":  BUDGET_COLORS.get(b.category, "bg-gray-100 text-gray-700"),
        })

    return JSONResponse(result)


# ── POST /api/budgets to update user budgets ─────────────────────────────────
@router.post("/api/budgets")
def update_user_budgets(
    payload: dict = Body(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    for category, amount in payload.items():
        budget = db.query(Budget).filter(
            Budget.user_id == current_user.id,
            Budget.category == category
        ).first()
        
        if budget:
            budget.budget_amount = float(amount)
        else:
            new_budget = Budget(
                user_id=current_user.id,
                category=category,
                budget_amount=float(amount)
            )
            db.add(new_budget)
            
    db.commit()
    return JSONResponse({"status": "success", "message": "Budgets updated successfully"})


# ── GET Financial Statement (daily, weekly, monthly, yearly, custom) ───────────
@router.get("/api/transactions/statement")
def get_statement(
    period: str = "monthly",  # daily, weekly, monthly, yearly, custom
    month: Optional[int] = None,
    year: Optional[int] = None,
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    now = datetime.datetime.utcnow()
    
    if month and year:
        target_year = year
        target_month = month
        start_date = datetime.datetime(target_year, target_month, 1, 0, 0, 0)
        if target_month == 12:
            end_date = datetime.datetime(target_year + 1, 1, 1, 0, 0, 0) - datetime.timedelta(seconds=1)
        else:
            end_date = datetime.datetime(target_year, target_month + 1, 1, 0, 0, 0) - datetime.timedelta(seconds=1)
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"
    elif period == "daily":
        start_date = datetime.datetime(now.year, now.month, now.day)
        end_date = start_date + datetime.timedelta(days=1) - datetime.timedelta(seconds=1)
        date_range_label = now.strftime("%d %b %Y")
    elif period == "weekly":
        start_date = now - datetime.timedelta(days=7)
        end_date = now
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {now.strftime('%d %b %Y')}"
    elif period == "yearly":
        start_date = datetime.datetime(now.year, 1, 1)
        end_date = datetime.datetime(now.year, 12, 31, 23, 59, 59)
        date_range_label = f"01 Jan {now.year} – 31 Dec {now.year}"
    elif period == "custom" and start_date_str and end_date_str:
        try:
            start_date = datetime.datetime.fromisoformat(start_date_str.replace("Z", ""))
            end_date = datetime.datetime.fromisoformat(end_date_str.replace("Z", ""))
            end_date = datetime.datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)
            date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"
        except Exception:
            start_date = datetime.datetime(now.year, now.month, 1)
            end_date = now
            date_range_label = f"01 {now.strftime('%b %Y')} – {now.strftime('%d %b %Y')}"
    else:  # monthly
        start_date = datetime.datetime(now.year, now.month, 1)
        if now.month == 12:
            end_date = datetime.datetime(now.year + 1, 1, 1) - datetime.timedelta(seconds=1)
        else:
            end_date = datetime.datetime(now.year, now.month + 1, 1) - datetime.timedelta(seconds=1)
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"

    txns = db.query(Transaction).filter(
        Transaction.user_id == current_user.id,
        Transaction.when >= start_date,
        Transaction.when <= end_date
    ).order_by(Transaction.when.desc()).all()

    total_income = sum([abs(t.amount) for t in txns if t.amount > 0 or t.payment_status == "credit"])
    total_expense = sum([abs(t.amount) for t in txns if t.amount < 0 or t.payment_status == "debit"])

    # Day-wise Grouped Breakdown
    grouped_days = {}
    for t in txns:
        day_key = t.when.strftime("%Y-%m-%d")
        if day_key not in grouped_days:
            diff = (now.date() - t.when.date()).days
            if diff == 0:
                display = "Today (" + t.when.strftime("%d %b %Y") + ")"
            elif diff == 1:
                display = "Yesterday (" + t.when.strftime("%d %b %Y") + ")"
            else:
                display = t.when.strftime("%A, %d %b %Y")

            grouped_days[day_key] = {
                "date": day_key,
                "display_date": display,
                "day_income": 0.0,
                "day_expense": 0.0,
                "transactions": []
            }
        
        amt = abs(t.amount)
        if t.amount > 0 or t.payment_status == "credit":
            grouped_days[day_key]["day_income"] += amt
        else:
            grouped_days[day_key]["day_expense"] += amt

        serialized = _serialize_txn(t)
        serialized["amount"] = amt
        grouped_days[day_key]["transactions"].append(serialized)

    day_wise_breakdown = list(grouped_days.values())

    return JSONResponse({
        "user_name": current_user.name,
        "email": current_user.email,
        "period": period.capitalize(),
        "date_range": date_range_label,
        "generated_at": now.strftime("%d %b %Y, %I:%M %p"),
        "total_income": round(total_income, 2),
        "total_expense": round(total_expense, 2),
        "net_savings": round(total_income - total_expense, 2),
        "transactions_count": len(txns),
        "day_wise_breakdown": day_wise_breakdown,
        "transactions": [_serialize_txn(t) for t in txns]
    })


@router.get("/api/transactions/statement/excel")
def get_statement_excel(
    period: str = "monthly",
    month: Optional[int] = None,
    year: Optional[int] = None,
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    now = datetime.datetime.utcnow()
    if month and year:
        target_year = year
        target_month = month
        start_date = datetime.datetime(target_year, target_month, 1, 0, 0, 0)
        if target_month == 12:
            end_date = datetime.datetime(target_year + 1, 1, 1, 0, 0, 0) - datetime.timedelta(seconds=1)
        else:
            end_date = datetime.datetime(target_year, target_month + 1, 1, 0, 0, 0) - datetime.timedelta(seconds=1)
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"
    elif period == "daily":
        start_date = datetime.datetime(now.year, now.month, now.day)
        end_date = start_date + datetime.timedelta(days=1) - datetime.timedelta(seconds=1)
        date_range_label = now.strftime("%d %b %Y")
    elif period == "weekly":
        start_date = now - datetime.timedelta(days=7)
        end_date = now
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {now.strftime('%d %b %Y')}"
    elif period == "yearly":
        start_date = datetime.datetime(now.year, 1, 1)
        end_date = datetime.datetime(now.year, 12, 31, 23, 59, 59)
        date_range_label = f"01 Jan {now.year} – 31 Dec {now.year}"
    elif period == "custom" and start_date_str and end_date_str:
        try:
            start_date = datetime.datetime.fromisoformat(start_date_str.replace("Z", ""))
            end_date = datetime.datetime.fromisoformat(end_date_str.replace("Z", ""))
            end_date = datetime.datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)
            date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"
        except Exception:
            start_date = datetime.datetime(now.year, now.month, 1)
            end_date = now
            date_range_label = f"01 {now.strftime('%b %Y')} – {now.strftime('%d %b %Y')}"
    else:  # monthly
        start_date = datetime.datetime(now.year, now.month, 1)
        if now.month == 12:
            end_date = datetime.datetime(now.year + 1, 1, 1) - datetime.timedelta(seconds=1)
        else:
            end_date = datetime.datetime(now.year, now.month + 1, 1) - datetime.timedelta(seconds=1)
        date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"

    txns = db.query(Transaction).filter(
        Transaction.user_id == current_user.id,
        Transaction.when >= start_date,
        Transaction.when <= end_date
    ).order_by(Transaction.when.desc()).all()

    total_income = sum([abs(t.amount) for t in txns if t.amount > 0 or t.payment_status == "credit"])
    total_expense = sum([abs(t.amount) for t in txns if t.amount < 0 or t.payment_status == "debit"])
    net_savings = total_income - total_expense

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Statement"

    # Styling definitions
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    table_header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    table_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # App Title Header
    ws.merge_cells('A1:F1')
    ws['A1'] = "FIM — Financial Intelligence Manager Statement"
    ws['A1'].font = title_font
    ws['A1'].fill = header_fill
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Summary Info
    ws['A3'] = "Account Name:"
    ws['B3'] = current_user.name
    ws['A4'] = "Email:"
    ws['B4'] = current_user.email
    ws['A5'] = "Period Range:"
    ws['B5'] = date_range_label
    ws['A6'] = "Generated Date:"
    ws['B6'] = now.strftime("%d %b %Y, %I:%M %p")

    ws['D3'] = "Total Income:"
    ws['E3'] = round(total_income, 2)
    ws['D4'] = "Total Expense:"
    ws['E4'] = round(total_expense, 2)
    ws['D5'] = "Net Savings:"
    ws['E5'] = round(net_savings, 2)

    for r in range(3, 7):
        ws[f'A{r}'].font = bold_font
        ws[f'D{r}'].font = bold_font

    # Table Column Headers
    headers = ["Date & Time", "Description / Title", "Category", "Type", "Amount (₹)", "Payment Status"]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_num)
        cell.value = header_title
        cell.font = table_header_font
        cell.fill = table_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[8].height = 26

    # Rows Data
    row_idx = 9
    for t in txns:
        is_credit = t.amount > 0 or t.payment_status == "credit"
        amt = abs(t.amount)
        ws.cell(row=row_idx, column=1, value=t.when.strftime("%Y-%m-%d %H:%M")).border = thin_border
        ws.cell(row=row_idx, column=2, value=t.name).border = thin_border
        ws.cell(row=row_idx, column=3, value=t.category).border = thin_border
        ws.cell(row=row_idx, column=4, value="Credit" if is_credit else "Debit").border = thin_border
        
        amt_cell = ws.cell(row=row_idx, column=5, value=amt)
        amt_cell.border = thin_border
        amt_cell.number_format = '#,##0.00'
        if is_credit:
            amt_cell.font = Font(color="16A34A", bold=True)
        else:
            amt_cell.font = Font(color="DC2626")

        ws.cell(row=row_idx, column=6, value=t.payment_status or ("credit" if is_credit else "debit")).border = thin_border
        row_idx += 1

    # Auto-fit columns width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"FIM_Statement_{period}_{current_user.name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )



